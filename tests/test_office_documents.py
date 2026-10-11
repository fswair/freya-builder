from __future__ import annotations

import hashlib
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.comments import Comment
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from freya_hackathon.builder import KIT, load_inputs
from freya_hackathon.input_documents import decode_document


def pdf_bytes(*texts: str) -> bytes:
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(612, 792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = stream
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_word_paragraph_table_order_and_footer_are_preserved(tmp_path):
    document = Document()
    document.add_paragraph("Önce kimlik doğrulaması yapılır.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Koşul"
    table.cell(0, 1).text = "İşlem"
    table.cell(1, 0).text = "Üçüncü kişi"
    table.cell(1, 1).text = "Bilgi paylaşılmaz."
    document.add_paragraph("Son olarak görüşme kapatılır.")
    document.sections[0].footer.paragraphs[0].text = "Ek kural: yalnız iş günleri."
    path = tmp_path / "policy.docx"
    document.save(path)
    text = decode_document(path, path.read_bytes())
    assert text.index("Önce") < text.index("Koşul") < text.index("Son olarak")
    assert '["Üçüncü kişi", "Bilgi paylaşılmaz."]' in text
    assert "Ek kural: yalnız iş günleri." in text


def test_excel_retains_all_sheets_dates_formulas_formats_and_notes(tmp_path):
    book = Workbook()
    first = book.active
    first.title = "Müşteriler"
    first.append(["Kimlik", "Tarih", "Tutar", "Toplam"])
    first.append([47, datetime(2026, 1, 5), 12.5, "=C2*2"])
    first["A2"].number_format = "000000"
    first["C2"].comment = Comment("Örnek müşteri tutarı.", "Fixture")
    other = book.create_sheet("Ek koşul")
    other.sheet_state = "hidden"
    other["A1"] = "Gizli sekme de girdinin parçasıdır."
    path = tmp_path / "customers.xlsx"
    book.save(path)
    text = decode_document(path, path.read_bytes())
    for term in ["Müşteriler", "2026-01-05", "000000", '"formula": "=C2*2"',
                 '"cached_value": null', "Örnek müşteri tutarı", "visibility: hidden",
                 "Gizli sekme de girdinin parçasıdır."]:
        assert term in text
    assert '"calculated_by_builder": false' in text


def test_pdf_keeps_page_boundaries_and_all_text():
    text = decode_document(Path("conditions.pdf"), pdf_bytes("First rule", "Second rule"))
    assert "[Page 1]\nFirst rule" in text
    assert "[Page 2]\nSecond rule" in text


def test_pdf_without_extractable_page_text_stops_instead_of_dropping_it():
    with pytest.raises(ValueError, match="page 2 has no extractable text"):
        decode_document(Path("scanned.pdf"), pdf_bytes("Readable page", ""))


def test_mixed_document_directory_reaches_builder_with_original_hashes(tmp_path):
    (tmp_path / "01-brief.md").write_text("Read every supplied document.")
    doc = Document()
    doc.add_paragraph("Kimlik kontrolü zorunlu.")
    doc.save(tmp_path / "02-policy.docx")
    (tmp_path / "03-rules.pdf").write_bytes(pdf_bytes("Close every call."))
    book = Workbook()
    book.active.append(["id", "amount"])
    book.active.append(["example-01", 125.5])
    book.save(tmp_path / "04-customers.xlsx")
    inputs = load_inputs([tmp_path], KIT / "tools/openapi.json")
    assert len(inputs.sources) == 5
    assert "Kimlik kontrolü zorunlu." in inputs.sources["brief-2"]
    assert "Close every call." in inputs.sources["brief-3"]
    assert "example-01" in inputs.sources["brief-4"]
    for record in inputs.manifest[:4]:
        assert record["sha256"] == hashlib.sha256(Path(record["path"]).read_bytes()).hexdigest()
