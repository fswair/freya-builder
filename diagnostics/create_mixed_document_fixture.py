"""Create synthetic, non-final input documents for a generic builder integration smoke."""

from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def create(output: Path) -> None:
    output.mkdir(parents=True, exist_ok=False)
    (output / "BRIEF.md").write_text(
        "# Synthetic development fixture: service desk reference lookup\n"
        "This is not a hackathon final case. Build a short text agent for reference lookup. "
        "Policy.docx contains every dialogue rule, Service-contract.md specifies the tools, "
        "Constraints.pdf adds honesty rules, and Examples.xlsx contains synthetic mock data. "
        "The agent must get results from tools; never embed the example records in prompts/code.\n",
        encoding="utf-8",
    )
    (output / "Service-contract.md").write_text(
        "# Service contract\n"
        "Base URL: supplied separately at builder invocation. No authentication required for "
        "this isolated synthetic fixture. All requests/responses use application/json.\n"
        "POST /lookup: required body field reference (string). "
        'Success HTTP200: {"data":{"found":true,"status":"ready"}}. '
        'Unknown reference HTTP200: {"data":{"found":false}}. '
        'Server failure HTTP503: {"error":"unavailable"}. '
        "This is a read-only lookup; there are no other service operations.\n",
        encoding="utf-8",
    )
    doc = Document()
    doc.add_heading("Dialogue requirements", level=1)
    doc.add_paragraph("Conduct the entire customer conversation in Spanish.")
    doc.add_paragraph(
        "Introduce yourself as an automated reference assistant and ask for a reference."
    )
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Condition"
    table.cell(0, 1).text = "Required response"
    for condition, response in [
        ("A reference is supplied", "Call POST /lookup with exactly that reference."),
        ("found is true", "Tell the returned status, say goodbye and end the conversation."),
        ("found is false", "Say the reference was not found, say goodbye and end."),
        ("Tool failure", "Say the service is unavailable, say goodbye and end; no retry."),
        ("Customer declines", "Say goodbye and end; do not call a tool."),
    ]:
        row = table.add_row()
        row.cells[0].text = condition
        row.cells[1].text = response
    doc.save(output / "Policy.docx")
    writer = PdfWriter()
    page = writer.add_blank_page(612, 792)
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })}),
    })
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 40 700 Td "
        b"(Never invent lookup results. Do not call any other service.) Tj ET"
    )
    page[NameObject("/Contents")] = stream
    buffer = BytesIO()
    writer.write(buffer)
    (output / "Constraints.pdf").write_bytes(buffer.getvalue())
    book = Workbook()
    book.active.title = "Synthetic examples"
    book.active.append(["reference", "found", "status"])
    book.active.append(["fixture-reference-71", True, "ready"])
    book.active.append(["fixture-reference-92", False, ""])
    book.save(output / "Examples.xlsx")
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("out", type=Path)
    create(parser.parse_args().out)
