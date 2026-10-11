"""Read supplied bank documents without executing macros, formulas or external links."""

from __future__ import annotations

import json
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

MAX_EXTRACTED_CHARS = 1_000_000
MAX_ARCHIVE_BYTES = 80_000_000
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _check_archive(data: bytes) -> None:
    with ZipFile(BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_ARCHIVE_BYTES:
            raise ValueError("Office archive expands beyond the document limit")


def _docx(data: bytes) -> str:
    from docx import Document
    from docx.table import Table

    _check_archive(data)
    document = Document(BytesIO(data))
    lines: list[str] = []

    def blocks(container) -> list[str]:
        result = []
        for block in container.iter_inner_content():
            if isinstance(block, Table):
                result.append("[Table]")
                for index, row in enumerate(block.rows, 1):
                    # JSON keeps line breaks and column boundaries unambiguous.
                    cells = ["\n".join(blocks(cell)) for cell in row.cells]
                    result.append(f"row {index}: " + json.dumps(cells, ensure_ascii=False))
                result.append("[/Table]")
            elif block.text.strip():
                result.append(block.text)
        return result

    lines.extend(blocks(document))
    # Preserve policy text outside the main body, including notes and text boxes.
    # Main-body text boxes are omitted by python-docx's Paragraph.text.
    with ZipFile(BytesIO(data)) as archive:
        for name in sorted(archive.namelist()):
            if name == "word/document.xml":
                root = ElementTree.fromstring(archive.read(name))
                for box in root.iter(W + "txbxContent"):
                    texts = ["".join(p.itertext()) for p in box.iter(W + "p")]
                    lines.append("[Text box]\n" + "\n".join(texts))
            elif (
                name.startswith("word/")
                and name.endswith(".xml")
                and (
                    Path(name).name.startswith(("header", "footer"))
                    or Path(name).name in {"footnotes.xml", "endnotes.xml", "comments.xml"}
                )
            ):
                root = ElementTree.fromstring(archive.read(name))
                texts = ["".join(t.text or "" for t in p.iter(W + "t")) for p in root.iter(W + "p")]
                if any(text.strip() for text in texts):
                    lines.append(f"[{name}]\n" + "\n".join(texts))
    return "\n".join(lines)


def _cell_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    _check_archive(data)
    book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    cached = load_workbook(BytesIO(data), data_only=True, keep_links=False)
    lines = []
    try:
        for sheet in book.worksheets:
            if sheet.max_row * sheet.max_column > 200_000:
                raise ValueError(f"Sheet {sheet.title!r} exceeds the cell limit")
            lines.append(f"[Sheet: {sheet.title}; visibility: {sheet.sheet_state}]")
            if sheet.merged_cells.ranges:
                lines.append("Merged cells: " + ", ".join(sorted(map(str, sheet.merged_cells))))
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    if cell.value is None and not cell.comment:
                        continue
                    value = {"cell": cell.coordinate, "value": _cell_value(cell.value)}
                    if cell.number_format != "General":
                        value["number_format"] = cell.number_format
                    if cell.data_type == "f":
                        value["formula"] = value.pop("value")
                        cached_value = cached[sheet.title][cell.coordinate].value
                        value["cached_value"] = _cell_value(cached_value)
                        value["calculated_by_builder"] = False
                    if cell.comment:
                        value["comment"] = cell.comment.text
                    if sheet.row_dimensions[cell.row].hidden:
                        value["hidden_row"] = True
                    values.append(value)
                if values:
                    lines.append(json.dumps(values, ensure_ascii=False))
        return "\n".join(lines)
    finally:
        book.close()
        cached.close()


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("Encrypted PDF requires an unlocked copy")
    pages = []
    for number, page in enumerate(reader.pages, 1):
        contents = page.get_contents()
        if contents and len(contents.get_data()) > 20_000_000:
            raise ValueError(f"PDF page {number} exceeds the content-stream limit")
        text = page.extract_text(extraction_mode="layout", layout_mode_strip_rotated=False)
        if not text.strip():
            # A blank extracted page may contain a scanned policy. Never discard it.
            raise ValueError(f"PDF page {number} has no extractable text; supply OCR text")
        pages.append(f"[Page {number}]\n{text.strip()}")
    return "\n\n".join(pages)


def extract_office_document(path: Path, data: bytes) -> str:
    extractor = {".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pdf": _pdf}
    try:
        text = extractor[path.suffix.lower()](data)
    except Exception as exc:
        raise ValueError(f"Cannot extract {path.name}: {exc}. No document was skipped.") from exc
    if not text.strip():
        raise ValueError(f"No text extracted from {path.name}. No document was skipped.")
    if len(text) > MAX_EXTRACTED_CHARS:
        raise ValueError(f"Extracted text exceeds the limit: {path.name}; no text was truncated")
    return text
